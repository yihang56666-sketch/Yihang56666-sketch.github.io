# -*- coding: utf-8 -*-
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PRIVATE = ROOT / "docs" / "private"
ACCENT = RGBColor(0x1D, 0x35, 0x57)
TEXT = RGBColor(0x24, 0x2C, 0x33)
MUTED = RGBColor(0x5B, 0x65, 0x70)

SOURCES = [
    PRIVATE / "Yihang-双选会简历-2026-final.md",
    PRIVATE / "四项目深剖-2026-final.md",
    PRIVATE / "面试题与讲稿-2026-final.md",
]

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
LIST_RE = re.compile(r"^(\d+)\.\s+(.*)$")


def set_font(run, *, size=10.5, bold=False, color=None, name="Segoe UI", east_asia="Microsoft YaHei"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)


def add_hyperlink(paragraph, label: str, url: str, size=10.5):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1D3557")
    rPr.append(color)
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Segoe UI")
    rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rPr.append(rFonts)
    run.append(rPr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str, *, size=10.5, color=TEXT):
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            set_font(run, size=size, bold=True, color=color)
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            set_font(run, size=size - 0.5, color=RGBColor(0x33, 0x40, 0x4A), name="Consolas", east_asia="Microsoft YaHei")
        else:
            match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", chunk)
            if match:
                add_hyperlink(paragraph, match.group(1), match.group(2), size=size)
            else:
                run = paragraph.add_run(chunk)
                set_font(run, size=size, color=color)


def add_bottom_rule(paragraph, color="1D3557"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def configure_styles(document: Document):
    normal = document.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(5)

    title = document.styles["Title"]
    title.font.name = "Segoe UI"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = ACCENT
    title.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    for name, size in (("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 12)):
        style = document.styles[name]
        style.font.name = "Segoe UI"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(5)


def add_page_footer(document: Document, title: str):
    section = document.sections[0]
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    set_font(run, size=8, color=MUTED)
    run = paragraph.add_run(" · Page ")
    set_font(run, size=8, color=MUTED)
    for kind, text in (("begin", None), ("instr", " PAGE "), ("separate", None), ("cached", "1"), ("end", None)):
        if kind == "instr":
            run = paragraph.add_run()
            set_font(run, size=8, color=MUTED)
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._element.append(instr)
        elif kind in {"begin", "separate", "end"}:
            run = paragraph.add_run()
            set_font(run, size=8, color=MUTED)
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._element.append(fld)
        else:
            run = paragraph.add_run(text)
            set_font(run, size=8, color=MUTED)


def add_toc(document: Document):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    for kind, text in (("begin", None), ("instr", ' TOC \\o "1-2" \\h \\z \\u '), ("separate", None), ("end", None)):
        run = paragraph.add_run()
        set_font(run, size=10, color=MUTED)
        if kind == "instr":
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._element.append(instr)
        elif kind in {"begin", "separate", "end"}:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._element.append(fld)
        else:
            placeholder = OxmlElement("w:t")
            placeholder.text = "目录将在 Word 打开后更新。"
            run._element.append(placeholder)


def enable_update_fields(document: Document):
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def add_table(document: Document, rows: list[list[str]]):
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(rows):
        for col_index, cell_text in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            add_inline(cell.paragraphs[0], cell_text, size=9.5)
            if row_index == 0:
                cell.paragraphs[0].runs and set_font(
                    cell.paragraphs[0].runs[0], size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF)
                )
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "1D3557")
                tcPr.append(shd)


def render_markdown(source: Path, target: Path):
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    title = source.stem.replace("-final", "")
    add_page_footer(document, "Yihang · " + title)
    enable_update_fields(document)

    lines = source.read_text(encoding="utf-8").splitlines()
    toc_added = False
    h1_seen = False
    after_intro = False
    is_resume = "简历" in source.name
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line:
            continue

        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            add_inline(paragraph, line[2:].strip(), size=24, color=ACCENT)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(4)
            h1_seen = True
            after_intro = False
            continue

        if line.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline(paragraph, line[3:].strip(), size=17, color=ACCENT)
            add_bottom_rule(paragraph)
            if h1_seen and not after_intro and not is_resume and not toc_added:
                add_toc(document)
                toc_added = True
                after_intro = True
            continue

        if line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline(paragraph, line[4:].strip(), size=13.5, color=ACCENT)
            continue

        if line.startswith("#### "):
            paragraph = document.add_paragraph(style="Heading 3")
            add_inline(paragraph, line[5:].strip(), size=12, color=ACCENT)
            continue

        if line.strip() == "---":
            continue

        if line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.4)
            add_inline(paragraph, line[2:].strip(), size=10, color=MUTED)
            paragraph.paragraph_format.space_after = Pt(6)
            continue

        if line.startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].startswith("|"):
                row = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in row):
                    rows.append(row)
                i += 1
            add_table(document, rows)
            continue

        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = 1.15
            add_inline(paragraph, line[2:].strip())
            continue

        numbered = LIST_RE.match(line)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = 1.15
            add_inline(paragraph, numbered.group(2))
            continue

        paragraph = document.add_paragraph()
        add_inline(paragraph, line.strip())

    document.save(target)
    print(f"Wrote {target}")


def main():
    for source in SOURCES:
        render_markdown(source, source.with_suffix(".docx"))


if __name__ == "__main__":
    main()
