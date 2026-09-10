# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()

# page margins
for section in doc.sections:
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

def set_run_font(run, size=10.5, bold=False, color=None, name_cn="微软雅黑", name_en="Segoe UI"):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = name_en
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name_cn)
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, 18, True, (40, 32, 60))
    p.paragraph_format.space_after = Pt(4)
    return p

def add_meta(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, 9.5, False, (90, 80, 110))
    p.paragraph_format.space_after = Pt(10)
    return p

def add_h2(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, 12, True, (110, 90, 180))
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    # bottom border
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pBdr.makeelement(qn("w:bottom"), {
        qn("w:val"): "single",
        qn("w:sz"): "6",
        qn("w:space"): "2",
        qn("w:color"): "B3A4F0",
    })
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def add_para(text, size=10, bold=False, indent=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size, bold)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if indent:
        p.paragraph_format.left_indent = Cm(0.2)
    return p

def add_bullet(text, size=9.5):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.1
    return p

def add_project(name, tech, repo, bullets):
    p = doc.add_paragraph()
    run = p.add_run(name)
    set_run_font(run, 11, True, (30, 25, 45))
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(1)
    p2 = doc.add_paragraph()
    run2 = p2.add_run(f"{tech}  |  {repo}")
    set_run_font(run2, 9, False, (100, 90, 120))
    p2.paragraph_format.space_after = Pt(2)
    for b in bullets:
        add_bullet(b)

add_title("Yihang")
add_meta("求职方向：嵌入式软件 / 工具链 / 客户端开发 / AI 应用工程")
add_meta("邮箱 yihang56666@gmail.com  ·  GitHub Yihang56666-sketch  ·  博客 yihang56666-sketch.github.io")

add_h2("能力概述")
for t in [
    "能独立把复杂需求拆成可演示、可回归、可讲清边界的完整作品：架构、状态模型、测试与安全门控。",
    "跨端交付：React + TypeScript 客户端、Electron / Capacitor Android、Python 工具链、无构建静态站。",
    "工程可靠性：单元测试、契约测试、类型检查、发布边界测试；默认安全、默认不吞错误。",
    "调试叙事：能讲清根因、取舍与回归，而不是只报功能列表。",
]:
    add_bullet(t, 9.5)

add_h2("项目经历")

add_project(
    "Hardware Butler · 安全优先的嵌入式开发助手",
    "Python 3.10+ · CLI + PyQt6 · 9 阶段工作流 · LLM 多 Provider",
    "github.com/yihang56666-sketch/hardware-butler",
    [
        "一句话需求编排为：需求解析→芯片选型→资料→CubeMX→固件→构建→烧录→观测→验证；状态机可 resume。",
        "默认 mock；真实烧录需环境变量 opt-in + 确认 token + 值域防呆 + 固件 hash + 审计日志。",
        "覆盖 14 个 MCU 厂商族；5 层行为验证（关键词/频率/文本/正则/数值范围）。",
        "LLM 只产结构化意图，执行器与门控由确定性代码掌控。",
        "回归：1017 passed / 12 skipped；ruff + mypy 全绿；GUI 13 tab、CLI 35 命令。",
    ],
)

add_project(
    "BEID / RIXIA · 本地优先学习工作台",
    "React 19 · TypeScript · Vite 7 · Zustand · Electron · Capacitor Android · PWA",
    "github.com/yihang56666-sketch/RIXIA",
    [
        "任务/习惯/日记/专注/考研 + B 站搜索、DASH 播放、时间点笔记同一领域模型；数据默认只落本机。",
        "自研 MSE 管线：sidx 索引 seek、背压与缓冲淘汰、CDN 备用源；不嵌官方 iframe。",
        "Zustand persist schema v3 + v1/v2 迁移；配额失败可见；companion backup。",
        "修复真实缺陷：BV 直链误判、逾期任务导航丢失、跨午夜 UTC 日、专注钳制不一致。",
        "回归：762 Vitest + 32 Node；typecheck 与生产构建通过。",
    ],
)

add_project(
    "magent · Codex 原生子智能体编排 Skill",
    "声明式 Skill 协议 · 零运行时依赖 · 标准库 unittest 契约测试",
    "github.com/yihang56666-sketch/magent",
    [
        "定位是编排契约层，不是自研运行时：原生工具跑子智能体，Skill 约束派谁/读什么/交什么证据。",
        "五个协作失败模式产品化：授权门、阻塞留本地、内容指纹基线、超时≠失败、释放并发槽。",
        "33 项无依赖结构契约测试；公开声明结构测试≠模型行为保证。",
        "与 LangGraph/CrewAI 分层：通用框架做执行引擎，本项目做判断、边界与核验层。",
    ],
)

add_project(
    "beid blog · 无构建静态作品集博客",
    "原生 HTML/CSS/JS · Hash SPA · GitHub Pages · Playwright",
    "yihang56666-sketch.github.io",
    [
        "零构建零框架：Hash 路由、动态 SEO 元数据、404 回退、深浅主题、无障碍与 reduced-motion。",
        "16 条 Playwright E2E：移动端溢出、主题持久化、中文 IME、TOC 锚点、动效降级。",
        "双选会前诚实化：真实字数与 Atom feed、本地 vendor、统一内外页视觉语言。",
    ],
)

add_h2("技术栈")
add_para("语言与框架：TypeScript、React 19、Python、JavaScript、HTML / CSS", 9.5)
add_para("客户端与站点：Vite、Electron、Capacitor Android、PWA、GitHub Pages、PyQt6", 9.5)
add_para("状态与数据：Zustand 持久化、schema 迁移、JSON 备份导入导出", 9.5)
add_para("测试与质量：Vitest、Playwright、pytest、unittest、ruff、mypy、契约测试、发布边界测试", 9.5)
add_para("领域：嵌入式工作流编排、安全执行门控、B 站 DASH/MSE、本地优先应用、多智能体协作协议", 9.5)

add_h2("工程判断")
for t in [
    "先把风险边界讲清楚，再谈自动化程度。",
    "能用测试固定的行为，不靠口头记忆。",
    "默认安全、默认可回滚、默认不吞错误。",
    "能演示的直接演示；不能演示的主动说明验证范围。",
]:
    add_bullet(t, 9.5)

add_h2("90 秒演示路线")
add_para("1. 博客：首页工程定位 → 项目卡 → 暗色主题 → 源仓库入口", 9.5)
add_para("2. BEID：今日下一步 → 逾期任务仍可见 → 本地备份 → 播放器联动", 9.5)
add_para("3. Hardware Butler：fixture mock 九阶段，说明默认不碰硬件与真实烧录门控", 9.5)
add_para("4. magent：现场跑 33 项契约测试，说明只读与主智能体写入权", 9.5)

add_h2("诚实边界")
for t in [
    "Hardware Butler 真实板卡烧录未在本材料宣称完成；有 runbook 与门控设计。",
    "BEID 扫码登录、真机播放与 Windows 安装包需现场设备验收。",
    "magent 契约测试验证协议结构，不承诺模型必然遵循或并行一定提速。",
    "博客 Hash 路由不是 SSR，SEO 依赖元数据同步。",
]:
    add_bullet(t, 9)

add_para("")
add_para("教育经历：现场投递时补充（学校 / 专业 / 学历 / 时间）", 9.5)
add_para("本稿只写入可复现的项目事实与当前实测数字。", 9, False)

out = r"D:\boke\docs\Yihang-双选会简历-2026.docx"
doc.save(out)
print("saved", out)
